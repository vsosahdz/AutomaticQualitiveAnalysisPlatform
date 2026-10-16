"""Importación de la codificación humana desde CSV o XLSX.

Formato esperado (columnas, sin distinguir mayúsculas): ``document``,
``quotation``, ``codes``, ``comment``. El campo ``codes`` admite varios códigos
separados por coma o punto y coma. Cada (cita, código) se vuelve una entrada.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import List, Optional

_CODE_SEP = re.compile(r"[;,]")


def _strip_accents(text: str) -> str:
    """Quita acentos para comparar de forma insensible a tildes."""
    return "".join(
        c for c in unicodedata.normalize("NFD", text)
        if unicodedata.category(c) != "Mn"
    )


def normalize_code(raw: str) -> Optional[str]:
    """Normaliza un código humano hacia el codebook canónico.

    Reglas (según la convención de codificación humana):
    - vacío → ``None`` (se ignora),
    - contiene "factor" (factor/factores/…) → "Barrera",
    - parecido a facilitador (facilitad…) → "Facilitador",
    - parecido a estrategia (estrateg…) → "Estrategia",
    - parecido a barrera (barrer…) → "Barrera",
    - etiquetado pero sin coincidencia → se conserva tal cual (no se descarta).
    """
    s = _strip_accents((raw or "").strip().lower())
    if not s:
        return None
    if "factor" in s:
        return "Barrera"
    if "facilitad" in s:
        return "Facilitador"
    if "estrateg" in s:
        return "Estrategia"
    if "barrer" in s:
        return "Barrera"
    return raw.strip()


@dataclass
class HumanEntry:
    """Una entrada de la codificación humana: cita + un código + comentario."""

    document: str
    quotation: str
    code: str
    comment: str = ""


def _split_codes(value: str) -> List[str]:
    """Separa el campo de códigos en códigos individuales, sin vacíos."""
    return [c.strip() for c in _CODE_SEP.split(value or "") if c.strip()]


def _rows_to_entries(rows: List[dict]) -> List[HumanEntry]:
    """Convierte filas (dicts con claves normalizadas) en entradas por código."""
    entries: List[HumanEntry] = []
    for row in rows:
        quotation = (row.get("quotation") or "").strip()
        if not quotation:
            continue
        document = (row.get("document") or "").strip()
        comment = (row.get("comment") or "").strip()
        for raw in _split_codes(row.get("codes") or ""):
            code = normalize_code(raw)
            if not code:
                continue  # código en blanco → se ignora
            entries.append(
                HumanEntry(document=document, quotation=quotation, code=code, comment=comment)
            )
    return entries


def _normalize_keys(row: dict) -> dict:
    """Normaliza las claves de una fila a minúsculas y sin espacios."""
    return {(k or "").strip().lower(): v for k, v in row.items()}


def parse_csv(content: bytes) -> List[HumanEntry]:
    """Parsea la codificación humana desde bytes de un CSV."""
    text = content.decode("utf-8-sig")  # tolera BOM
    reader = csv.DictReader(io.StringIO(text))
    rows = [_normalize_keys(r) for r in reader]
    return _rows_to_entries(rows)


def parse_xlsx(content: bytes) -> List[HumanEntry]:
    """Parsea la codificación humana desde bytes de un XLSX (primera hoja)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []
    rows: List[dict] = []
    for values in rows_iter:
        row = {header[i]: (values[i] if i < len(values) else None) for i in range(len(header))}
        # Los valores pueden no ser str (números); se normalizan a str más abajo.
        rows.append({k: ("" if v is None else str(v)) for k, v in row.items()})
    return _rows_to_entries(rows)


def parse_human_file(filename: str, content: bytes) -> List[HumanEntry]:
    """Parsea la codificación humana según la extensión del archivo.

    Args:
        filename: nombre del archivo (para detectar CSV vs XLSX).
        content: bytes del archivo.

    Returns:
        Lista de :class:`HumanEntry` (una por cita y código).

    Raises:
        ValueError: si la extensión no es soportada.
    """
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv(content)
    if lower.endswith(".xlsx"):
        return parse_xlsx(content)
    raise ValueError("Formato no soportado: usa CSV o XLSX")
