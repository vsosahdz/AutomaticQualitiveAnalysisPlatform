"""Generación del reporte DOCX de hallazgos solo-IA.

Construye un documento Word con, por cada documento comparado: una explicación
en lenguaje natural del reporte léxico (interpretando P/R/F1 y Cohen's kappa) y la
lista de hallazgos que solo encontró la IA (con su cita, código, justificación,
similitud y la relevancia marcada por la persona).
"""
from __future__ import annotations

import io

from docx import Document as Docx
from docx.shared import Pt, RGBColor

from app.models.report import ReportDocument, ReportRequest


def _interpret_kappa(k: float) -> str:
    """Interpreta el valor de kappa según la escala de Landis & Koch."""
    if k < 0:
        return "sin acuerdo (peor que el azar)"
    if k < 0.20:
        return "acuerdo leve"
    if k < 0.40:
        return "acuerdo aceptable"
    if k < 0.60:
        return "acuerdo moderado"
    if k < 0.80:
        return "acuerdo sustancial"
    return "acuerdo casi perfecto"


def _pct(x: float) -> str:
    """Formatea una fracción 0..1 como porcentaje entero."""
    return f"{round(x * 100)}%"


def _lexical_explanation(doc: ReportDocument) -> str:
    """Redacta en lenguaje natural la interpretación del reporte léxico."""
    o = doc.overall
    kappa_txt = (
        f" El acuerdo corregido por azar (Cohen's kappa) fue de {o.kappa:.2f}, "
        f"lo que indica un {_interpret_kappa(o.kappa)}."
        if o.kappa is not None
        else ""
    )
    return (
        f"Tomando la codificación humana como referencia, el LLM alcanzó una "
        f"precisión de {_pct(o.precision)} (de todo lo que el LLM marcó, esa "
        f"proporción coincidió con la persona), un recall de {_pct(o.recall)} "
        f"(de todo lo que la persona marcó, esa proporción también la identificó "
        f"el LLM) y un F1 de {_pct(o.f1)} (media armónica de precisión y recall)."
        f"{kappa_txt} Estas medidas son léxicas: exigen coincidir en la misma "
        f"oración y el mismo código. Por eso, los hallazgos que aparecen abajo —"
        f"detectados por similitud semántica— capturan aportes del LLM que no "
        f"coinciden literalmente con la codificación humana y que conviene revisar."
    )


def _add_metrics_table(docx: Docx, doc: ReportDocument) -> None:
    """Agrega la tabla de métricas por código (con nombres completos)."""
    headers = ["Código", "Precisión (P)", "Recall (R)", "F1", "Kappa (κ)", "TP", "FP", "FN"]
    table = docx.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for m in doc.per_code:
        cells = table.add_row().cells
        vals = [m.code, f"{m.precision:.2f}", f"{m.recall:.2f}", f"{m.f1:.2f}",
                "—" if m.kappa is None else f"{m.kappa:.2f}", str(m.tp), str(m.fp), str(m.fn)]
        for i, v in enumerate(vals):
            cells[i].text = v
    # Fila global.
    o = doc.overall
    cells = table.add_row().cells
    gvals = ["GLOBAL", f"{o.precision:.2f}", f"{o.recall:.2f}", f"{o.f1:.2f}",
             "—" if o.kappa is None else f"{o.kappa:.2f}", "", "", ""]
    for i, v in enumerate(gvals):
        cells[i].text = v


def _add_context_with_bold_quote(paragraph, context: str) -> None:
    """Agrega el contexto a un párrafo, poniendo en negritas la cita marcada «…».

    El frontend envía la cita rodeada de «…»; aquí se resalta ese tramo en
    negritas (sin mostrar los delimitadores) para que se identifique fácilmente.
    """
    i = context.find("«")
    j = context.find("»", i + 1) if i != -1 else -1
    if i == -1 or j == -1:
        paragraph.add_run(context)
        return
    paragraph.add_run(context[:i])
    paragraph.add_run(context[i + 1 : j]).bold = True
    paragraph.add_run(context[j + 1 :])


def _relevance_label(value: str) -> str:
    """Etiqueta legible de la relevancia marcada."""
    return {"relevant": "Relevante", "no": "No relevante"}.get(value, "Sin marcar")


def build_docx(request: ReportRequest) -> bytes:
    """Construye el reporte DOCX y devuelve sus bytes.

    Args:
        request: datos del reporte (documentos, métricas y hallazgos solo-IA).

    Returns:
        Los bytes del archivo .docx.
    """
    docx = Docx()
    docx.add_heading("Reporte de hallazgos del LLM (complementariedad)", level=0)
    intro = docx.add_paragraph()
    intro.add_run(
        "Este reporte resume, por documento, el acuerdo entre la codificación "
        "humana y la del LLM, y destaca los hallazgos que solo encontró el LLM. "
    )
    intro.add_run(
        f"La similitud semántica se calculó con el modelo de embeddings "
        f"'{request.embedding_model}' y un umbral de {request.threshold:.2f} para "
        f"considerar que dos citas expresan el mismo hallazgo."
    )

    # Definiciones de los códigos, una sola vez al inicio.
    if request.definitions:
        docx.add_heading("Definiciones de los códigos", level=1)
        for label, definition in request.definitions.items():
            if not definition:
                continue
            p = docx.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(definition)

    for doc in request.documents:
        docx.add_heading(doc.name, level=1)

        docx.add_heading("Interpretación del reporte léxico", level=2)
        docx.add_paragraph(_lexical_explanation(doc))
        if doc.diff_counts is not None:
            dc = doc.diff_counts
            docx.add_paragraph(
                f"Distribución del acuerdo por oración: {dc.agree} en acuerdo, "
                f"{dc.only_llm} solo del LLM, {dc.only_human} solo humanas y "
                f"{dc.partial} parciales."
            )
        _add_metrics_table(docx, doc)

        docx.add_heading(f"Hallazgos solo-IA ({len(doc.only_ai)})", level=2)
        if not doc.only_ai:
            docx.add_paragraph("No se detectaron hallazgos solo-IA bajo el umbral actual.")
        for idx, f in enumerate(doc.only_ai, start=1):
            p = docx.add_paragraph()
            run = p.add_run(f"{idx}. [{f.code}] ")
            run.bold = True
            p.add_run(f"“{f.quote}”")
            meta = docx.add_paragraph()
            meta_run = meta.add_run(
                f"Similitud máx. con lo humano: {f.max_similarity:.2f} · "
                f"Relevancia: {_relevance_label(f.relevancia)}"
            )
            meta_run.italic = True
            meta_run.font.size = Pt(9)
            meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            # Contexto del documento con la cita en NEGRITAS (marcada con «…»).
            if f.context:
                cp = docx.add_paragraph()
                cp.add_run("Contexto en el documento: ").bold = True
                _add_context_with_bold_quote(cp, f.context)
            if f.justification:
                jp = docx.add_paragraph()
                jp.add_run("Justificación del LLM: ").bold = True
                jp.add_run(f.justification)
            if f.best_match_quote:
                docx.add_paragraph(
                    f"Cita humana más parecida: “{f.best_match_quote}” "
                    f"[{f.best_match_code}]"
                )

    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()
